# !/usr/bin/env python3
# -*- coding:utf-8 -*-

# @Time    : 2024/3/18 17:34
# @Author  : wangchongshi
# @Email   : wangchongshi.wcs@antgroup.com
# @FileName: pptx_reader.py
from pathlib import Path
from typing import Dict, Iterable, Iterator, List, Optional, Union

from agentuniverse.agent.action.knowledge.reader.reader import Reader
from agentuniverse.agent.action.knowledge.store.document import Document


class PptxReader(Reader):
    """Read text shapes and tables, including those inside nested groups.

    Each text shape or non-empty table produces one document in shape order.
    Table cells are separated by tabs and rows by newlines; spanned cells are
    blank so merged content is emitted only at its origin.
    """

    def _load_data(self, file: Union[str, Path], ext_info: Optional[Dict] = None) -> List[Document]:
        """Parse the pptx file.

        Note:
            `python-pptx` is required to read PPTX files: `pip install python-pptx`
        """
        try:
            from pptx import Presentation
        except ImportError:
            raise ImportError(
                "python-pptx is required to read pptx files: `pip install python-pptx`"
            )
        if isinstance(file, str):
            file = Path(file)
        presentation = Presentation(file)
        document_list = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            for text in self._iter_shape_text(slide.shapes):
                metadata = {"slide_number": slide_number, "file_name": file.name}
                if ext_info is not None:
                    metadata.update(ext_info)
                document_list.append(Document(text=text, metadata=metadata))
        return document_list

    def _iter_shape_text(self, shapes: Iterable) -> Iterator[str]:
        """Yield text from shape containers in their stored order."""
        from pptx.enum.shapes import MSO_SHAPE_TYPE

        for shape in shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                yield from self._iter_shape_text(shape.shapes)
            elif shape.has_table:
                text = "\n".join(
                    "\t".join("" if cell.is_spanned else cell.text
                              for cell in row.cells)
                    for row in shape.table.rows
                )
                if text.strip():
                    yield text
            elif hasattr(shape, "text"):
                yield shape.text
