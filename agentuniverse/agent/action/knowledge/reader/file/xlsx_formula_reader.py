"""Read formula text alongside stored XLSX results without evaluating formulas."""
import hashlib
import json
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from pydantic import StrictInt

from agentuniverse.agent.action.knowledge.reader.reader import Reader
from agentuniverse.agent.action.knowledge.store.document import Document


class XlsxFormulaReader(Reader):
    """Opt-in dual-view reader; original XlsxReader behavior remains unchanged.

    Results are cached workbook values, possibly stale, never recalculated.
    Missing cached values are explicitly marked, not represented as zero.
    """
    max_file_bytes: StrictInt = 16 * 1024 * 1024
    max_uncompressed_bytes: StrictInt = 128 * 1024 * 1024
    max_cells: StrictInt = 100000

    def _validate_options(self):
        for name in ("max_file_bytes", "max_uncompressed_bytes", "max_cells"):
            if type(getattr(self, name)) is not int or getattr(self, name) <= 0:
                raise ValueError(f"{name} must be a positive integer")

    def _load_data(self, file, ext_info=None):
        self._validate_options()
        import openpyxl
        if isinstance(file, (str, Path)):
            name = Path(file).name
            with open(file, "rb") as source:
                content = source.read(self.max_file_bytes + 1)
        else:
            position = file.tell()
            try:
                file.seek(0)
                content = file.read(self.max_file_bytes + 1)
            finally:
                file.seek(position)
            name_value = getattr(file, "name", "unknown")
            name = Path(name_value).name if isinstance(name_value, (str, Path)) else "unknown"
        if not isinstance(content, bytes) or len(content) > self.max_file_bytes:
            raise ValueError("expected XLSX bytes within max_file_bytes")
        with ZipFile(BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > self.max_uncompressed_bytes:
                raise ValueError("workbook exceeds max_uncompressed_bytes")
        digest = hashlib.sha256(content).hexdigest()
        formulas = None
        cached = None
        try:
            formulas = openpyxl.load_workbook(BytesIO(content), data_only=False, read_only=True, keep_links=False)
            cached = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True, keep_links=False)
            total = sum((sheet.max_row or 0) * (sheet.max_column or 0) for sheet in formulas.worksheets)
            if total > self.max_cells:
                raise ValueError("worksheet dimensions exceed max_cells")
            result = []
            visited = 0
            for sheet in formulas.worksheets:
                # Paired workbooks come from the exact same bytes, including sheet order.
                values = cached[sheet.title]
                # Exporters can underreport the used range. Keep the conservative
                # declared-size preflight above, but stream the actual rows in
                # both views instead of silently dropping out-of-range formulas.
                sheet.reset_dimensions()
                values.reset_dimensions()
                rows = []
                annotations = []
                for formula_row, value_row in zip(sheet.iter_rows(), values.iter_rows()):
                    # With unknown width, openpyxl emits empty tuples for missing
                    # rows. Charge those too so sparse row gaps remain bounded.
                    if not formula_row:
                        visited += 1
                        if visited > self.max_cells:
                            raise ValueError("workbook exceeds max_cells")
                    cells = []
                    for cell, value_cell in zip(formula_row, value_row):
                        visited += 1
                        if visited > self.max_cells:
                            raise ValueError("workbook exceeds max_cells")
                        value = value_cell.value
                        if cell.data_type == "f":
                            if not isinstance(cell.value, str):
                                raise ValueError("array/data-table formulas are not supported")
                            status = "missing" if value is None else "stored_unverified"
                            # Scalars become text deliberately for portable JSON metadata;
                            # Excel errors remain tagged by their original cell type.
                            annotation = {"cell": cell.coordinate, "formula": cell.value,
                                          "cached_value": None if value is None else str(value),
                                          "cached_type": value_cell.data_type, "status": status}
                            annotations.append(annotation)
                            cells.append(json.dumps(annotation, ensure_ascii=False))
                        else:
                            cells.append("" if value is None else str(value))
                    if any(cells):
                        rows.append(" | ".join(cells))
                if rows:
                    metadata = dict(ext_info or {})
                    metadata.update({"file_name": name, "sheet_name": sheet.title,
                                     "workbook_sha256": digest, "formula_cells": annotations,
                                     "formula_results_verified": False})
                    text = "\n".join(rows)
                    identity = hashlib.sha256((digest + "\0" + sheet.title).encode("utf-8")).hexdigest()
                    result.append(Document(id=identity, text=text, metadata=metadata))
            return result
        finally:
            if formulas is not None:
                formulas.close()
            if cached is not None:
                cached.close()

    def _initialize_by_component_configer(self, config):
        names = ("max_file_bytes", "max_uncompressed_bytes", "max_cells")
        validated = XlsxFormulaReader(**{name: getattr(config, name, getattr(self, name)) for name in names})
        validated._validate_options()
        super()._initialize_by_component_configer(config)
        for name in names:
            setattr(self, name, getattr(validated, name))
        return self
