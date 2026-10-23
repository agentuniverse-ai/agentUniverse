"""Unsupported legacy formats must not receive mislabeled OOXML output."""
import asyncio
import json

import pytest
from openpyxl import load_workbook

from agentuniverse.agent.action.tool.common_tool.excel_tool import ExcelTool


@pytest.mark.parametrize('suffix', ['.xls', '.XLS', '.xlsb', '.XLSB'])
@pytest.mark.parametrize('asynchronous', [False, True])
def test_unsupported_write_does_not_create_file(tmp_path, suffix, asynchronous):
    path = tmp_path / ('new' + suffix)
    tool = ExcelTool(allowed_extensions=['.xlsx', '.xls', '.xlsb'])
    kwargs = {'file_path': str(path), 'mode': 'write', 'data': [['value'], [42]]}
    result = json.loads(asyncio.run(tool.async_execute(**kwargs)) if asynchronous
                        else tool.execute(**kwargs))
    assert result['status'] == 'error'
    assert 'convert' in result['error'].lower()
    assert not path.exists()


@pytest.mark.parametrize('mode', ['write', 'read', 'append', 'info'])
@pytest.mark.parametrize('suffix', ['.xls', '.xlsb'])
def test_unsupported_operations_preserve_existing_bytes(tmp_path, mode, suffix):
    path = tmp_path / ('existing' + suffix)
    original = b'existing legacy file placeholder'
    path.write_bytes(original)
    tool = ExcelTool(allowed_extensions=['.xlsx', '.xls', '.xlsb'])

    result = json.loads(tool.execute(str(path), mode, data=[['changed']], overwrite=True))
    assert result['status'] == 'error'
    assert 'convert' in result['error'].lower()
    assert path.read_bytes() == original


@pytest.mark.parametrize('suffix', ['.xlsx', '.XLSX'])
def test_supported_xlsx_roundtrip_retained(tmp_path, suffix):
    path = tmp_path / ('supported' + suffix)
    tool = ExcelTool()
    assert json.loads(tool.execute(str(path), 'write', data=[['Name'], ['first']]))['status'] == 'success'
    assert json.loads(tool.execute(str(path), 'append', data=[['second']]))['status'] == 'success'
    assert json.loads(tool.execute(str(path), 'read'))['status'] == 'success'
    assert json.loads(tool.execute(str(path), 'info'))['status'] == 'success'
    book = load_workbook(path)
    try:
        assert list(book.active.values) == [('Name',), ('first',), ('second',)]
    finally:
        book.close()


def test_custom_allowlist_still_restricts_supported_formats(tmp_path):
    path = tmp_path / 'excluded.xlsx'
    result = json.loads(ExcelTool(allowed_extensions=[]).execute(str(path), 'write', data=[[1]]))
    assert result['status'] == 'error'
    assert not path.exists()


def test_default_and_custom_allowlist_contract(tmp_path):
    assert ExcelTool().allowed_extensions == ['.xlsx', '.xls']
    tool = ExcelTool(allowed_extensions=['.xlsm', '.xltx', '.xltm'])
    for suffix in tool.allowed_extensions:
        assert tool._validate_path(str(tmp_path / ('custom' + suffix))) == {'valid': True}


@pytest.mark.parametrize('suffix', ['.xls', '.xlsb'])
@pytest.mark.parametrize('mode', ['read', 'write', 'append', 'info'])
@pytest.mark.parametrize('asynchronous', [False, True])
def test_custom_backend_can_handle_allowed_legacy_formats(tmp_path, suffix, mode, asynchronous):
    class CustomBackend(ExcelTool):
        def _read_excel(self, file_path, **kwargs):
            return {'status': 'success', 'backend': 'custom', 'operation': 'read'}

        def _write_excel(self, file_path, **kwargs):
            return {'status': 'success', 'backend': 'custom', 'operation': 'write'}

        def _append_excel(self, file_path, **kwargs):
            return {'status': 'success', 'backend': 'custom', 'operation': 'append'}

        def _get_excel_info(self, file_path):
            return {'status': 'success', 'backend': 'custom', 'operation': 'info'}

    # Preserve the inherited .xls allowlist as well as explicit .xlsb opt-in.
    tool = CustomBackend() if suffix == '.xls' else CustomBackend(allowed_extensions=[suffix])
    path = str(tmp_path / ('custom' + suffix))
    kwargs = {'file_path': path, 'mode': mode, 'data': [[1]]}
    result = json.loads(asyncio.run(tool.async_execute(**kwargs)) if asynchronous
                        else tool.execute(**kwargs))
    assert result == {'status': 'success', 'backend': 'custom', 'operation': mode}


@pytest.mark.parametrize('suffix', ['.xls', '.xlsb'])
def test_builtin_writer_rejects_legacy_path_before_creating_directory(tmp_path, suffix):
    path = tmp_path / 'not_created' / ('new' + suffix)
    # Direct backend calls are protected too; shared validation is not bypassed
    # to permit a mislabeled write.
    result = ExcelTool()._write_excel(str(path), [[1]])
    assert result['status'] == 'error'
    assert not path.parent.exists()


@pytest.mark.parametrize('suffix', ['.xls', '.xlsb'])
@pytest.mark.parametrize('mode', ['read', 'append', 'info'])
def test_renamed_ooxml_copy_is_not_legacy_support(tmp_path, suffix, mode):
    from shutil import copyfile
    from openpyxl import Workbook

    source = tmp_path / 'source.xlsx'
    workbook = Workbook()
    workbook.active.append(['original'])
    workbook.save(source)
    workbook.close()
    path = tmp_path / ('renamed' + suffix)
    copyfile(source, path)
    original = path.read_bytes()
    tool = ExcelTool(allowed_extensions=['.xlsx', '.xls', '.xlsb'])
    result = json.loads(tool.execute(str(path), mode, data=[['changed']]))
    assert result['status'] == 'error'
    assert path.read_bytes() == original == source.read_bytes()
