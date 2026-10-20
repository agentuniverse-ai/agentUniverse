"""Every public Excel read exit closes its actual read-only ZIP archive."""
import asyncio
import json
from unittest.mock import patch
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook

from agentuniverse.agent.action.tool.common_tool.excel_tool import ExcelTool


@pytest.mark.parametrize('case', ['success', 'bounded', 'missing_sheet', 'malformed_cell'])
@pytest.mark.parametrize('asynchronous', [False, True])
def test_read_closes_loaded_archive(tmp_path, case, asynchronous):
    path = tmp_path / 'input.xlsx'
    book = Workbook()
    book.active.append([1, 2])
    book.active.append([3, 4])
    book.save(path)
    book.close()
    if case == 'malformed_cell':
        with ZipFile(path) as source:
            entries = {name: source.read(name) for name in source.namelist()}
        entries['xl/worksheets/sheet1.xml'] = entries['xl/worksheets/sheet1.xml'].replace(b'<v>1</v>', b'<v>not-a-number</v>')
        with ZipFile(path, 'w') as target:
            for name, value in entries.items():
                target.writestr(name, value)
    loaded = []

    def recording_load(*args, **kwargs):
        workbook = load_workbook(*args, **kwargs)
        loaded.append(workbook)
        return workbook

    kwargs = {'mode': 'read'}
    if case == 'missing_sheet':
        kwargs['sheet_name'] = 'Missing'
    elif case == 'bounded':
        kwargs['max_rows'] = 1
    try:
        with patch('openpyxl.load_workbook', side_effect=recording_load):
            tool = ExcelTool()
            result = json.loads(asyncio.run(tool.async_execute(str(path), **kwargs)) if asynchronous else tool.execute(str(path), **kwargs))
        assert result['status'] == ('error' if case in {'missing_sheet', 'malformed_cell'} else 'success')
        assert len(loaded) == 1
        assert loaded[0]._archive.fp is None
    finally:
        # Ensure a failing baseline regression does not retain fixture resources.
        for workbook in loaded:
            workbook.close()


def test_workbook_load_failure_keeps_error_response(tmp_path):
    path = tmp_path / 'invalid.xlsx'
    path.write_bytes(b'not a zip archive')
    result = json.loads(ExcelTool().execute(str(path), mode='read'))
    assert result['status'] == 'error'
    assert 'Failed to read Excel file' in result['error']
