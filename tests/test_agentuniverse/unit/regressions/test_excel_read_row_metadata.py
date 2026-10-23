"""Read widths describe the returned rows even without worksheet dimensions."""
import json
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from agentuniverse.agent.action.tool.common_tool.excel_tool import ExcelTool


@pytest.mark.parametrize('max_rows,expected_width,expected_rows', [(None, 3, 2), (1, 1, 1), (2, 3, 2)])
def test_unsized_sheet_uses_widest_returned_row(tmp_path, max_rows, expected_width, expected_rows):
    path = tmp_path / 'unsized.xlsx'
    book = Workbook(write_only=True)
    sheet = book.create_sheet('Data')
    sheet.append(['short'])
    sheet.append(['wide', 'second', 'third'])
    book.save(path)
    book.close()
    with ZipFile(path) as archive:
        assert b'<dimension' not in archive.read('xl/worksheets/sheet1.xml')
    result = json.loads(ExcelTool().execute(str(path), mode='read', max_rows=max_rows))
    assert result['status'] == 'success'
    assert result['columns'] == expected_width
    assert result['rows_read'] == expected_rows
    assert result['data'][0] == ['short']
    # Unknown worksheet totals stay unknown; no unbounded dimension scan.
    assert result['total_rows'] is None
    assert result['total_columns'] is None


def test_empty_unsized_sheet_has_zero_returned_width(tmp_path):
    path = tmp_path / 'empty.xlsx'
    book = Workbook(write_only=True)
    book.create_sheet('Data')
    book.save(path)
    book.close()
    result = json.loads(ExcelTool().execute(str(path), mode='read'))
    assert result['status'] == 'success'
    assert result['columns'] == 0
    assert result['data'] == []


def test_sized_sheet_width_and_totals_remain_unchanged(tmp_path):
    path = tmp_path / 'sized.xlsx'
    book = Workbook()
    book.active.append(['short'])
    book.active.append(['wide', 'second', 'third'])
    book.save(path)
    book.close()
    result = json.loads(ExcelTool().execute(str(path), mode='read', max_rows=1))
    assert result['columns'] == 3
    assert result['data'] == [['short', None, None]]
    assert result['total_rows'] == 2
    assert result['total_columns'] == 3
