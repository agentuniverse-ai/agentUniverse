"""Reported append row ranges match the actual openpyxl insertion positions."""
import json

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from agentuniverse.agent.action.tool.common_tool.excel_tool import ExcelTool


@pytest.mark.parametrize("initial,expected_start", [("blank", 1), ("populated", 3), ("styled", 2)])
@pytest.mark.parametrize("rows", [[["first"], ["second"]], [[], ["middle"], []]])
def test_append_reports_actual_rows(tmp_path, initial, expected_start, rows):
    path = tmp_path / "rows.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    if initial == "populated":
        sheet.append(["header"])
        sheet.append(["existing"])
    elif initial == "styled":
        sheet["A1"].font = Font(bold=True)
    workbook.save(path)
    workbook.close()

    result = json.loads(ExcelTool().execute(str(path), mode="append", sheet_name="Data", data=rows))
    assert result["status"] == "success"
    assert result["rows_appended"] == len(rows)
    assert result["start_row"] == expected_start
    assert result["end_row"] == expected_start + len(rows) - 1
    reopened = load_workbook(path)
    try:
        sheet = reopened["Data"]
        for offset, row in enumerate(rows):
            if row:
                assert sheet.cell(expected_start + offset, 1).value == row[0]
        if initial == "populated":
            assert sheet["A1"].value == "header"
            assert sheet["A2"].value == "existing"
        elif initial == "styled":
            assert sheet["A1"].font.bold is True
    finally:
        reopened.close()


@pytest.mark.parametrize("rows,columns", [
    ([["short"], ["wide", "second", "third"]], 3),
    ([[], ["wide", "second"]], 2),
    ([["a", "b"], ["c", "d"]], 2),
    ([[], []], 0),
])
def test_write_reports_widest_input_row(tmp_path, rows, columns):
    path = tmp_path / "ragged.xlsx"
    result = json.loads(ExcelTool().execute(str(path), mode="write", data=rows))
    assert result["status"] == "success"
    assert result["columns"] == columns
    assert result["rows_written"] == len(rows)
    reopened = load_workbook(path)
    try:
        for row_index, row in enumerate(rows, 1):
            for column_index, value in enumerate(row, 1):
                assert reopened.active.cell(row_index, column_index).value == value
    finally:
        reopened.close()
