"""Read grouped PowerPoint content through real saved/reopened presentations."""

import pytest

pytest.importorskip("pptx")
from pptx import Presentation
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.util import Inches

from agentuniverse.agent.action.tool.common_tool.powerpoint_tool import PowerPointTool


def add_text(shapes, text):
    box = shapes.add_textbox(0, 0, Inches(2), Inches(1))
    box.text = text
    return box


def group_shapes(shapes, depth):
    for _ in range(depth):
        shapes = shapes.add_group_shape().shapes
    return shapes


def read(tmp_path, presentation, **limits):
    path = tmp_path / "grouped.pptx"
    presentation.save(path)
    result = PowerPointTool(base_dir=str(tmp_path), **limits).execute(mode="read", file_path=str(path))
    assert result["status"] == "success", result
    assert set(result["slides"][0]) == {"slide_number", "title", "texts", "tables", "notes"}
    return result


@pytest.mark.parametrize("depth", [0, 1, 3])
def test_grouped_text_keeps_depth_first_shape_order(tmp_path, depth):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_text(slide.shapes, "Before")
    shapes = group_shapes(slide.shapes, depth)
    add_text(shapes, "第一段\nSecond paragraph")
    add_text(shapes, "Next")
    add_text(slide.shapes, "After")
    result = read(tmp_path, presentation)
    assert result["slides"][0]["texts"] == ["Before", "第一段\nSecond paragraph", "Next", "After"]
    assert not result["truncated"]


@pytest.mark.parametrize("depth", [0, 1, 2])
def test_grouped_tables_keep_title_notes_and_table_cells(tmp_path, depth):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Report"
    shapes = group_shapes(slide.shapes, depth)
    add_text(shapes, "Before table")
    table_shape = slide.shapes.add_table(2, 2, 0, 0, Inches(2), Inches(1))
    for index, row in enumerate([["名称", "Value"], ["Revenue", "123"]]):
        for column, value in enumerate(row):
            table_shape.table.cell(index, column).text = value
    if depth:
        shapes._spTree.append(table_shape._element)
    add_text(shapes, "After table")
    slide.notes_slide.notes_text_frame.text = "Speaker notes"
    result = read(tmp_path, presentation)
    item = result["slides"][0]
    assert item["title"] == "Report"
    assert item["texts"] == ["Before table", "After table"]
    assert item["tables"] == [[["名称", "Value"], ["Revenue", "123"]]]
    assert item["notes"] == "Speaker notes"
    assert not result["truncated"]


def test_empty_groups_do_not_change_neighbor_texts(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_text(slide.shapes, "Before")
    group_shapes(slide.shapes, 4)
    add_text(slide.shapes, "After")
    result = read(tmp_path, presentation)
    assert result["slides"][0]["texts"] == ["Before", "After"]
    assert not result["truncated"]


@pytest.mark.parametrize("limit,expected", [(1, []), (2, ["A"]), (3, ["A", "B"]), (4, ["A", "B", "C"])])
def test_group_container_and_children_each_consume_shape_budget(tmp_path, limit, expected):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shapes = group_shapes(slide.shapes, 1)
    for value in ["A", "B", "C"]:
        add_text(shapes, value)
    result = read(tmp_path, presentation, max_shapes_per_slide=limit)
    assert result["slides"][0]["texts"] == expected
    assert result["truncated"] == (limit < 4)


def test_deep_groups_stop_at_shape_budget(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_text(group_shapes(slide.shapes, 8), "Deep")
    result = read(tmp_path, presentation, max_shapes_per_slide=3)
    assert result["slides"][0]["texts"] == []
    assert result["truncated"]


def test_empty_groups_are_bounded_too(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for _ in range(5):
        slide.shapes.add_group_shape()
    result = read(tmp_path, presentation, max_shapes_per_slide=2)
    assert result["slides"][0]["texts"] == []
    assert result["truncated"]


def test_group_table_keeps_row_and_column_budgets(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shapes = group_shapes(slide.shapes, 1)
    table_shape = slide.shapes.add_table(2, 2, 0, 0, Inches(2), Inches(1))
    for row in table_shape.table.rows:
        for cell in row.cells:
            cell.text = "cell"
    shapes._spTree.append(table_shape._element)
    result = read(tmp_path, presentation, max_table_rows=1, max_table_columns=1)
    assert result["slides"][0]["tables"] == [[["cell"]]]
    assert result["truncated"]


def test_group_text_keeps_character_budget(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_text(group_shapes(slide.shapes, 1), "Long grouped text")
    result = read(tmp_path, presentation, max_text_chars=5)
    assert result["slides"][0]["texts"] == ["Long…"]
    assert result["truncated"]


def test_shape_budget_resets_for_each_slide(tmp_path):
    presentation = Presentation()
    for text in ["First", "Second"]:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        add_text(group_shapes(slide.shapes, 1), text)
    result = read(tmp_path, presentation, max_shapes_per_slide=2)
    assert [slide["texts"] for slide in result["slides"]] == [["First"], ["Second"]]
    assert not result["truncated"]


@pytest.mark.parametrize("depth", [0, 1])
@pytest.mark.parametrize("title_last", [False, True])
def test_title_is_separate_from_exact_shape_budget(tmp_path, depth, title_last):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title = slide.shapes.title
    title.text = "Report"
    add_text(group_shapes(slide.shapes, depth), "Content")
    if title_last:
        slide.shapes._spTree.append(title._element)
    result = read(tmp_path, presentation, max_shapes_per_slide=depth + 1)
    assert result["slides"][0]["title"] == "Report"
    assert result["slides"][0]["texts"] == ["Content"]
    assert not result["truncated"]


@pytest.mark.parametrize("limit,expected", [(2, []), (3, ["Content"])])
def test_non_text_shapes_also_consume_group_budget(tmp_path, limit, expected):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shapes = group_shapes(slide.shapes, 1)
    shapes.add_connector(MSO_CONNECTOR.STRAIGHT, 0, 0, Inches(1), Inches(1))
    add_text(shapes, "Content")
    result = read(tmp_path, presentation, max_shapes_per_slide=limit)
    assert result["slides"][0]["texts"] == expected
    assert result["truncated"] == (limit < 3)


def test_info_retains_top_level_shape_counts(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shapes = group_shapes(slide.shapes, 2)
    add_text(shapes, "Content")
    path = tmp_path / "info.pptx"
    presentation.save(path)
    result = PowerPointTool(base_dir=str(tmp_path)).execute(mode="info", file_path=str(path))
    assert result["status"] == "success", result
    assert result["slides"][0]["shape_count"] == 1
    assert result["slides"][0]["table_count"] == 0
