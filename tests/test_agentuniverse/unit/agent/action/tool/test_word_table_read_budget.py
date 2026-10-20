"""Real-DOCX regressions for table text-budget boundaries."""

import pytest
from docx import Document

from agentuniverse.agent.action.tool.common_tool.word_document_tool import WordDocumentTool


def read_tables(tmp_path, rows, budget, *, paragraph=None, extra_table=False):
    """Write a DOCX fixture and read it through the public tool API."""
    document = Document()
    if paragraph is not None:
        document.add_paragraph(paragraph)
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for cells, values in zip(table.rows, rows, strict=True):
        for cell, value in zip(cells.cells, values, strict=True):
            cell.text = value
    if extra_table:
        document.add_table(rows=1, cols=1).cell(0, 0).text = "later"
    document.save(tmp_path / "table.docx")
    result = WordDocumentTool(base_dir=str(tmp_path), max_text_chars=budget).execute(
        mode="read", file_path="table.docx",
    )
    assert result["status"] == "success"
    return result


def test_exact_final_cell_budget_is_not_truncated(tmp_path):
    """Consuming exactly the budget is not itself omitted content."""
    result = read_tables(tmp_path, [["abc", "def"]], 6)
    assert result["tables"] == [[["abc", "def"]]]
    assert result["truncated"] is False


def test_exhausted_budget_stops_before_remaining_cells(tmp_path):
    """Do not emit fabricated empty cells once the text budget is consumed."""
    result = read_tables(tmp_path, [["abc", "def", "ghi"]], 3)
    assert result["tables"] == [[["abc"]]]
    assert result["truncated"] is True


def test_paragraph_exhaustion_does_not_emit_empty_table(tmp_path):
    """The table traversal shares the paragraph text budget."""
    result = read_tables(tmp_path, [["def", "ghi"]], 3, paragraph="abc")
    assert result["tables"] == []
    assert result["truncated"] is True


@pytest.mark.parametrize("rows,extra_table", [([["abc"], ["def"]], False), ([["abc"]], True)])
def test_exact_cell_boundary_still_reports_omitted_structure(tmp_path, rows, extra_table):
    """Later rows or tables must keep the truncation signal true."""
    result = read_tables(tmp_path, rows, 3, extra_table=extra_table)
    assert result["tables"] == [[["abc"]]]
    assert result["truncated"] is True


def test_mid_cell_cut_preserves_ellipsis(tmp_path):
    """Existing mid-field truncation formatting remains unchanged."""
    result = read_tables(tmp_path, [["abcdef"]], 3)
    assert result["tables"] == [[["ab…"]]]
    assert result["truncated"] is True


@pytest.mark.parametrize(
    "cells,expected,truncated",
    [(["", "abc"], ["", "abc"], False), (["abc", ""], ["abc"], True)],
)
def test_empty_cells_respect_structural_budget_boundary(tmp_path, cells, expected, truncated):
    """Keep visited empty cells, but do not traverse omitted ones after exhaustion."""
    result = read_tables(tmp_path, [cells], 3)
    assert result["tables"] == [[expected]]
    assert result["truncated"] is truncated


def test_exact_shared_paragraph_table_budget(tmp_path):
    """A table can consume the exact remainder after a paragraph without loss."""
    result = read_tables(tmp_path, [["c"]], 3, paragraph="ab")
    assert result["paragraphs"][0]["text"] == "ab"
    assert result["tables"] == [[["c"]]]
    assert result["truncated"] is False


def test_empty_table_remains_readable(tmp_path):
    """A table without rows consumes a block, not the text budget."""
    document = Document()
    document.add_table(rows=0, cols=1)
    document.save(tmp_path / "empty.docx")
    result = WordDocumentTool(base_dir=str(tmp_path), max_text_chars=3).execute(
        mode="read", file_path="empty.docx",
    )
    assert result["status"] == "success"
    assert result["tables"] == [[]]
    assert result["truncated"] is False
