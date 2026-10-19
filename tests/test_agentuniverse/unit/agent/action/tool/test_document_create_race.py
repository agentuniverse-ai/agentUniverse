"""Create-only document writes must survive a concurrent successful creator."""
import threading

from docx import Document
from docx.document import Document as WordDocument
from pptx import Presentation
from pptx.presentation import Presentation as PowerPointPresentation
import pytest

from agentuniverse.agent.action.tool.common_tool.powerpoint_tool import PowerPointTool
from agentuniverse.agent.action.tool.common_tool.word_document_tool import WordDocumentTool


@pytest.fixture(params=["word", "powerpoint"])
def document_tool(request, tmp_path):
    if request.param == "word":
        tool, document_class, extension = WordDocumentTool(base_dir=str(tmp_path)), WordDocument, ".docx"

        def payload(text):
            return {"blocks": [{"type": "paragraph", "text": text}]}

        def contents(path):
            return [paragraph.text for paragraph in Document(path).paragraphs]
    else:
        tool, document_class, extension = PowerPointTool(base_dir=str(tmp_path)), PowerPointPresentation, ".pptx"

        def payload(text):
            return {"slides": [{"layout": "title", "title": text}]}

        def contents(path):
            return [shape.text for slide in Presentation(path).slides
                    for shape in slide.shapes if shape.has_text_frame]
    return tool, document_class, extension, payload, contents


@pytest.mark.parametrize("overwrite", [False, True])
@pytest.mark.parametrize("with_template", [False, True])
def test_concurrent_creator_keeps_its_file_unless_overwrite_allowed(
        document_tool, tmp_path, monkeypatch, overwrite, with_template):
    tool, document_class, extension, payload, contents = document_tool
    target = tmp_path / ("result" + extension)
    options = {}
    if with_template:
        template = tmp_path / ("template" + extension)
        assert tool.execute(mode="create", file_path=str(template), **payload("template"))["status"] == "success"
        options["template_path"] = str(template)

    paused, release = threading.Event(), threading.Event()
    original_save = document_class.save
    first_result = []

    def first_create():
        first_result.append(tool.execute(mode="create", file_path=str(target),
                                         overwrite=overwrite, **payload("first"), **options))

    first = threading.Thread(target=first_create)

    def delayed_save(document, path):
        if threading.current_thread() is first:
            paused.set()
            assert release.wait(15)
        return original_save(document, path)

    # The native SDK still serializes both real documents; this only controls
    # the interleaving after the first caller's existence check.
    monkeypatch.setattr(document_class, "save", delayed_save)
    first.start()
    try:
        assert paused.wait(15)
        second = tool.execute(mode="create", file_path=str(target), **payload("winner"))
        assert second["status"] == "success"
        winner_bytes = target.read_bytes()
    finally:
        release.set()
        first.join(15)
    assert not first.is_alive()
    if overwrite:
        assert "first" in contents(target)
        assert "winner" not in contents(target)
        if with_template:
            assert "template" in contents(target)
    else:
        assert "winner" in contents(target)
        assert "first" not in contents(target)
        assert target.read_bytes() == winner_bytes
    assert first_result[0]["status"] == ("success" if overwrite else "error")
    assert not list(tmp_path.glob(".word-*"))
    assert not list(tmp_path.glob(".powerpoint-*"))


def test_already_existing_file_stays_unchanged(document_tool, tmp_path):
    tool, _, extension, payload, _ = document_tool
    target = tmp_path / ("existing" + extension)
    assert tool.execute(mode="create", file_path=str(target), **payload("winner"))["status"] == "success"
    before = target.read_bytes()
    result = tool.execute(mode="create", file_path=str(target), **payload("other"))
    assert result["status"] == "error"
    assert target.read_bytes() == before


def test_append_still_updates_the_existing_document(document_tool, tmp_path):
    tool, _, extension, payload, contents = document_tool
    target = tmp_path / ("append" + extension)
    assert tool.execute(mode="create", file_path=str(target), **payload("initial"))["status"] == "success"
    assert tool.execute(mode="append", file_path=str(target), **payload("appended"))["status"] == "success"
    assert "initial" in contents(target)
    assert "appended" in contents(target)
