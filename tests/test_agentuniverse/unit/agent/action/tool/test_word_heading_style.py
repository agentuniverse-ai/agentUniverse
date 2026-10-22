"""Regression coverage for explicit styles on Word heading blocks."""

from docx import Document
from docx.enum.style import WD_STYLE_TYPE

from agentuniverse.agent.action.tool.common_tool.word_document_tool import WordDocumentTool


def test_heading_honors_explicit_style(tmp_path):
    """An explicit heading style is retained in the saved DOCX."""
    tool = WordDocumentTool(base_dir=str(tmp_path))
    result = tool.execute(
        mode="create", file_path="heading.docx",
        blocks=[{"type": "heading", "text": "Title", "level": 1, "style": "Subtitle"}],
    )
    assert result["status"] == "success"
    assert Document(tmp_path / "heading.docx").paragraphs[0].style.name == "Subtitle"


def test_heading_default_still_uses_level(tmp_path):
    """Omitting the style retains the existing heading-level behavior."""
    tool = WordDocumentTool(base_dir=str(tmp_path))
    result = tool.execute(
        mode="create", file_path="heading.docx",
        blocks=[{"type": "heading", "text": "Title", "level": 2}],
    )
    assert result["status"] == "success"
    assert Document(tmp_path / "heading.docx").paragraphs[0].style.name == "Heading 2"


def test_heading_uses_custom_template_style_on_append(tmp_path):
    """Append supports a named paragraph style supplied by the document."""
    path = tmp_path / "heading.docx"
    document = Document()
    document.styles.add_style("Report Heading", WD_STYLE_TYPE.PARAGRAPH)
    document.save(path)
    tool = WordDocumentTool(base_dir=str(tmp_path))
    result = tool.execute(
        mode="append", file_path="heading.docx",
        blocks=[{"type": "heading", "text": "Title", "style": "Report Heading"}],
    )
    assert result["status"] == "success"
    assert Document(path).paragraphs[0].style.name == "Report Heading"


def test_unknown_heading_style_does_not_overwrite_existing_file(tmp_path):
    """Invalid styles follow paragraph-style errors without changing the file."""
    path = tmp_path / "heading.docx"
    Document().save(path)
    before = path.read_bytes()
    tool = WordDocumentTool(base_dir=str(tmp_path))
    result = tool.execute(
        mode="append", file_path="heading.docx",
        blocks=[{"type": "heading", "text": "Title", "style": "Missing style"}],
    )
    assert result["status"] == "error"
    assert path.read_bytes() == before


def test_heading_uses_custom_style_from_create_template(tmp_path):
    """Create resolves the requested heading style against the template."""
    template = Document()
    template.styles.add_style("Report Heading", WD_STYLE_TYPE.PARAGRAPH)
    template.save(tmp_path / "template.docx")
    tool = WordDocumentTool(base_dir=str(tmp_path))
    result = tool.execute(
        mode="create", file_path="heading.docx", template_path="template.docx",
        blocks=[{"type": "heading", "text": "Title", "style": "Report Heading"}],
    )
    assert result["status"] == "success"
    assert Document(tmp_path / "heading.docx").paragraphs[0].style.name == "Report Heading"
