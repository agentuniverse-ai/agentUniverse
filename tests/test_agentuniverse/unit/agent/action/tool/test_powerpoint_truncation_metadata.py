"""Read truncation metadata from actual saved PowerPoint files."""

import pytest
from pptx import Presentation

from agentuniverse.agent.action.tool.common_tool.powerpoint_tool import PowerPointTool


def read_deck(tmp_path, slides, limit):
    presentation = Presentation()
    for title, notes in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    path = tmp_path / "budget.pptx"
    presentation.save(path)
    result = PowerPointTool(base_dir=str(tmp_path), max_text_chars=limit).execute(
        mode="read", file_path=str(path))
    assert result["status"] == "success", result
    emitted = sum(len(slide["title"]) + len(slide["notes"])
                  + sum(map(len, slide["texts"])) for slide in result["slides"])
    assert emitted <= limit
    return result


@pytest.mark.parametrize("field", ["title", "notes"])
@pytest.mark.parametrize("limit,expected,cut", [(1, "…", True), (5, "abcd…", True),
                                               (11, "abcdefghijk", False), (12, "abcdefghijk", False)])
def test_title_and_notes_report_actual_text_cuts(tmp_path, field, limit, expected, cut):
    values = {"title": "", "notes": ""}
    values[field] = "abcdefghijk"
    result = read_deck(tmp_path, [(values["title"], values["notes"])], limit)
    assert result["slides"][0][field] == expected
    assert result["truncated"] is cut


@pytest.mark.parametrize("slides", [[("12345", ""), ("lost", "")],
                                    [("", "12345"), ("", "lost")]])
def test_exact_exhaustion_reports_skipped_later_slide(tmp_path, slides):
    result = read_deck(tmp_path, slides, 5)
    assert result["slide_count"] == 2
    assert len(result["slides"]) == 1
    assert result["truncated"] is True


def test_title_and_notes_share_budget_and_report_omitted_notes(tmp_path):
    result = read_deck(tmp_path, [("12345", "lost")], 5)
    assert result["slides"][0]["title"] == "12345"
    assert result["slides"][0]["notes"] == ""
    assert result["truncated"] is True


def test_multiple_slides_within_budget_remain_complete(tmp_path):
    result = read_deck(tmp_path, [("One", "A"), ("Two", "B")], 8)
    assert len(result["slides"]) == 2
    assert result["truncated"] is False
