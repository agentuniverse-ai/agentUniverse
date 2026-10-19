"""Exercise container text through real saved PowerPoint files."""

import pytest

from agentuniverse.agent.action.knowledge.reader.file.file_reader import FileReader
from agentuniverse.agent.action.knowledge.reader.file.pptx_reader import PptxReader


pptx = pytest.importorskip("pptx")
from pptx.util import Inches


def add_text(shapes, text):
    box = shapes.add_textbox(0, 0, Inches(2), Inches(1))
    box.text = text
    return box


def save_presentation(tmp_path, presentation):
    path = tmp_path / "content.pptx"
    presentation.save(path)
    return path


@pytest.mark.parametrize("depth", [0, 1, 2])
def test_text_shapes_in_nested_groups_keep_document_order(tmp_path, depth):
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_text(slide.shapes, "Before")
    shapes = slide.shapes
    for _ in range(depth):
        shapes = shapes.add_group_shape().shapes
    add_text(shapes, "第一段\nSecond paragraph")
    add_text(shapes, "Next")
    add_text(slide.shapes, "After")

    documents = PptxReader().load_data(save_presentation(tmp_path, presentation))

    assert [doc.text for doc in documents] == [
        "Before", "第一段\nSecond paragraph", "Next", "After"
    ]
    assert all(doc.metadata == {"slide_number": 1, "file_name": "content.pptx"}
               for doc in documents)


@pytest.mark.parametrize("depth", [0, 1, 2])
@pytest.mark.parametrize("through_file_reader", [False, True])
def test_tables_in_nested_groups_preserve_rows_and_cells(
    tmp_path, depth, through_file_reader
):
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    frame = slide.shapes.add_table(2, 2, 0, 0, Inches(4), Inches(2))
    values = [["产品", "数量"], ["苹果\nApple", "3"]]
    for row, cells in enumerate(values):
        for column, text in enumerate(cells):
            frame.table.cell(row, column).text = text
    for _ in range(depth):
        frame = slide.shapes.add_group_shape([frame])
    path = save_presentation(tmp_path, presentation)

    documents = (FileReader().load_data([path]) if through_file_reader
                 else PptxReader().load_data(path))

    assert [doc.text for doc in documents] == ["产品\t数量\n苹果\nApple\t3"]
    assert documents[0].metadata == {"slide_number": 1, "file_name": "content.pptx"}


def test_merged_table_cells_do_not_repeat_text(tmp_path):
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    table = slide.shapes.add_table(2, 2, 0, 0, Inches(4), Inches(2)).table
    table.cell(0, 0).text = "Heading"
    table.cell(0, 1).text = "Details"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = "Left"
    table.cell(1, 1).text = "Right"

    documents = PptxReader().load_data(save_presentation(tmp_path, presentation))

    assert [doc.text for doc in documents] == ["Heading\nDetails\t\nLeft\tRight"]


def test_containers_keep_slide_numbers_and_metadata_overrides(tmp_path):
    presentation = pptx.Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    add_text(slide.shapes.add_group_shape().shapes, "Grouped")
    table = slide.shapes.add_table(1, 1, 0, 0, Inches(2), Inches(1)).table
    table.cell(0, 0).text = "Tabular"
    path = save_presentation(tmp_path, presentation)
    extra = {"collection": "reports", "file_name": "upload.pptx"}

    documents = PptxReader().load_data(str(path), ext_info=extra)

    assert [doc.text for doc in documents] == ["Grouped", "Tabular"]
    assert all(doc.metadata == {
        "slide_number": 2, "file_name": "upload.pptx", "collection": "reports"
    } for doc in documents)
    documents[0].metadata["collection"] = "changed"
    assert documents[1].metadata["collection"] == "reports"
    assert extra == {"collection": "reports", "file_name": "upload.pptx"}


def test_empty_containers_do_not_add_documents(tmp_path):
    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_group_shape()
    slide.shapes.add_table(2, 2, 0, 0, Inches(4), Inches(2))

    assert PptxReader().load_data(save_presentation(tmp_path, presentation)) == []
