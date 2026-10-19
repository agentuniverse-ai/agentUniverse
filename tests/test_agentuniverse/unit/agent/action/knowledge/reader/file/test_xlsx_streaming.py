"""Read sparse real workbooks without creating a rectangular in-memory cell grid."""
from unittest.mock import patch
import pytest
openpyxl = pytest.importorskip('openpyxl')
from openpyxl.worksheet.worksheet import Worksheet
from agentuniverse.agent.action.knowledge.reader.file.xlsx_reader import XlsxReader


def test_sparse_workbook_streams_and_keeps_columns(tmp_path):
    path = tmp_path / 'sparse.xlsx'
    workbook = openpyxl.Workbook()
    workbook.active['A1'] = 'start'
    workbook.active.cell(2000, 100, 'end')
    workbook.save(path)
    workbook.close()
    original = openpyxl.load_workbook
    loaded = []
    def capture(*args, **kwargs):
        value = original(*args, **kwargs)
        loaded.append(value)
        return value
    with patch.object(openpyxl, 'load_workbook', side_effect=capture), \
         patch.object(Worksheet, 'cell', side_effect=AssertionError('must not materialize cells')):
        result = XlsxReader().load_data(path)
    assert loaded[0].read_only
    assert loaded[0]._archive.fp is None
    assert result[0].text.splitlines() == ['start' + ' | ' * 99, ' | ' * 99 + 'end']
    assert result[0].metadata['max_row'] == 2000
    assert result[0].metadata['max_col'] == 100


def test_workbook_closed_when_row_iteration_fails():
    from unittest.mock import MagicMock
    workbook = MagicMock()
    workbook.sheetnames = ['one']
    workbook['one'].iter_rows.side_effect = RuntimeError('read failed')
    with patch.object(openpyxl, 'load_workbook', return_value=workbook):
        with pytest.raises(RuntimeError, match='read failed'):
            XlsxReader().load_data('broken.xlsx')
    workbook.close.assert_called_once()


def test_multiple_sheets_values_and_metadata(tmp_path):
    path = tmp_path / 'values.xlsx'
    workbook = openpyxl.Workbook()
    workbook.active.append(['name', 0, False, None, '中文'])
    workbook.create_sheet('empty')
    workbook.create_sheet('second').append(['other'])
    workbook.save(path)
    workbook.close()
    result = XlsxReader().load_data(path, ext_info={'origin': 'test'})
    assert [doc.metadata['sheet_name'] for doc in result] == ['Sheet', 'second']
    assert result[0].text == 'name | 0 | False |  | 中文'
    assert result[1].text == 'other'
    assert all(doc.metadata['origin'] == 'test' for doc in result)
