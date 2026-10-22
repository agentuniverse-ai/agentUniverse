"""Chart-only Excel tabs must not discard readable worksheets."""

import pytest
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from agentuniverse.agent.action.knowledge.reader.file.file_reader import FileReader
from agentuniverse.agent.action.knowledge.reader.file.xlsx_reader import XlsxReader


def add_chart(workbook, data, index):
    chart = BarChart()
    chart.add_data(Reference(data, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    sheet = workbook.create_chartsheet("Chart", index=index)
    sheet.add_chart(chart)


@pytest.mark.parametrize("chart_index", [0, 1, 2])
@pytest.mark.parametrize("route", ["path", "string", "file_reader"])
def test_chartsheet_keeps_data_sheet_order_and_metadata(tmp_path, chart_index, route):
    workbook = Workbook()
    data = workbook.active
    data.title = "Data"
    data.append(["Name", "Value"])
    data.append(["A", 1])
    later = workbook.create_sheet("Later")
    later.append(["B", False])
    later.sheet_state = "hidden"
    add_chart(workbook, data, chart_index)
    workbook.create_sheet("Empty")
    path = tmp_path / "mixed.xlsx"
    workbook.save(path)
    extra = {"source": "fixture"}
    documents = (FileReader().load_data([path], ext_info=extra) if route == "file_reader"
                 else XlsxReader().load_data(str(path) if route == "string" else path, ext_info=extra))
    assert [doc.text for doc in documents] == ["Name | Value\nA | 1", "B | False"]
    assert [doc.metadata["sheet_name"] for doc in documents] == ["Data", "Later"]
    assert documents[0].metadata == {"file_name": "mixed.xlsx", "sheet_name": "Data",
                                     "max_row": 2, "max_col": 2, "source": "fixture"}
    assert extra == {"source": "fixture"}


def test_embedded_chart_does_not_exclude_its_data_worksheet(tmp_path):
    workbook = Workbook()
    data = workbook.active
    data.append([1, 2])
    chart = BarChart()
    chart.add_data(Reference(data, min_col=1, max_col=2, min_row=1))
    data.add_chart(chart, "D1")
    path = tmp_path / "embedded.xlsx"
    workbook.save(path)
    documents = XlsxReader().load_data(path)
    assert [doc.text for doc in documents] == ["1 | 2"]
    assert documents[0].metadata["sheet_name"] == "Sheet"
