"""Real XLSX archive fixtures with explicit cached formula values."""
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from unittest.mock import patch
from xml.etree import ElementTree

import openpyxl
import pytest

from agentuniverse.agent.action.knowledge.reader.file.xlsx_formula_reader import XlsxFormulaReader


def workbook(cached=None):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Amounts"
    sheet.append(["alpha", "total"])
    sheet.append([2, "=A2+3"])
    stream = BytesIO()
    book.save(stream)
    book.close()
    if cached is None:
        return stream.getvalue()
    output = BytesIO()
    with ZipFile(BytesIO(stream.getvalue())) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(content)
                ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                cell = root.find(".//s:c[@r='B2']/s:v", ns)
                assert cell is not None
                cell.text = cached
                content = ElementTree.tostring(root, encoding="utf-8")
            target.writestr(item, content)
    return output.getvalue()


@pytest.mark.parametrize("cached", [None, "0", "5"])
def test_formula_and_cached_value_distinguished(tmp_path, cached):
    path = tmp_path / "formulas.xlsx"
    path.write_bytes(workbook(cached))
    result = XlsxFormulaReader().load_data(path)
    annotation = result[0].metadata["formula_cells"][0]
    assert annotation["cell"] == "B2"
    assert annotation["formula"] == "=A2+3"
    assert annotation["cached_value"] == cached
    assert annotation["status"] == ("missing" if cached is None else "stored_unverified")
    assert "=A2+3" in result[0].text
    assert result[0].metadata["formula_results_verified"] is False


def test_stream_position_and_reserved_metadata():
    stream = BytesIO(workbook())
    stream.seek(10)
    metadata = {"source": "local", "formula_results_verified": True}
    result = XlsxFormulaReader().load_data(stream, ext_info=metadata)
    assert stream.tell() == 10
    assert not stream.closed
    assert result[0].metadata["formula_results_verified"] is False
    assert metadata["formula_results_verified"] is True


@pytest.mark.parametrize("options", [{"max_cells": 1}, {"max_file_bytes": 10}, {"max_uncompressed_bytes": 10}])
def test_resource_limits(options):
    with pytest.raises(ValueError):
        XlsxFormulaReader(**options).load_data(BytesIO(workbook()))


def test_deterministic_ids_and_input_file_unchanged(tmp_path):
    path = tmp_path / "formulas.xlsx"
    original = workbook("5")
    path.write_bytes(original)
    reader = XlsxFormulaReader()
    first = reader.load_data(path)
    second = reader.load_data(path)
    assert first == second
    assert path.read_bytes() == original


def test_knowledge_file_reader_override(tmp_path):
    from agentuniverse.agent.action.knowledge import knowledge as module
    path = tmp_path / "formulas.xlsx"
    path.write_bytes(workbook("5"))
    knowledge = module.Knowledge(readers={"xlsx": "formulas"})
    try:
        with patch.object(module, "ReaderManager") as manager:
            manager.return_value.get_instance_obj.return_value = XlsxFormulaReader()
            documents = knowledge._load_data(source_path=str(path))
            manager.return_value.get_instance_obj.assert_called_once_with("formulas")
        assert documents[0].metadata["formula_cells"][0]["cached_value"] == "5"
    finally:
        knowledge.insert_executor.shutdown(wait=True)
        knowledge.query_executor.shutdown(wait=True)


@pytest.mark.parametrize("declared", ["A1:A1", "A1:B1"])
def test_underreported_sheet_dimensions_do_not_hide_formulas(declared):
    output = BytesIO()
    with ZipFile(BytesIO(workbook("5"))) as source, ZipFile(output, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            content = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(content)
                dimension = root.find("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}dimension")
                dimension.set("ref", declared)
                content = ElementTree.tostring(root, encoding="utf-8")
            target.writestr(item, content)
    content = output.getvalue()
    documents = XlsxFormulaReader().load_data(BytesIO(content))
    assert documents[0].metadata["formula_cells"][0]["cell"] == "B2"
    assert documents[0].metadata["formula_cells"][0]["cached_value"] == "5"
    with pytest.raises(ValueError, match="max_cells"):
        XlsxFormulaReader(max_cells=3).load_data(BytesIO(content))


def test_dimensionless_empty_rows_are_charged_to_cell_budget():
    book = openpyxl.Workbook(write_only=True)
    sheet = book.create_sheet("Sparse")
    for _ in range(4):
        sheet.append([])
    sheet.append(["=1+1"])
    stream = BytesIO()
    book.save(stream)
    book.close()
    result = XlsxFormulaReader().load_data(BytesIO(stream.getvalue()))
    assert result[0].metadata["formula_cells"][0]["cell"] == "A5"
    with pytest.raises(ValueError, match="max_cells"):
        XlsxFormulaReader(max_cells=3).load_data(BytesIO(stream.getvalue()))


def test_real_registration_public_insert_persistence_and_query(tmp_path):
    from pathlib import Path
    import agentuniverse.base.annotation.trace as trace_module
    from agentuniverse.agent.action.knowledge.knowledge import Knowledge
    from agentuniverse.agent.action.knowledge.reader.reader_manager import ReaderManager
    from agentuniverse.agent.action.knowledge.store.store_manager import StoreManager
    from agentuniverse.agent.action.knowledge.doc_processor.doc_processor_manager import DocProcessorManager
    from agentuniverse.agent.action.knowledge.doc_processor.jieba_keyword_extractor import JiebaKeywordExtractor
    from agentuniverse.agent.action.knowledge.rag_router.rag_router_manager import RagRouterManager
    from agentuniverse.agent.action.knowledge.rag_router.base_router import BaseRouter
    from agentuniverse.agent.action.knowledge.store.sqlite_store import SQLiteStore
    from agentuniverse.base.config.configer import Configer
    from agentuniverse.base.config.component_configer.component_configer import ComponentConfiger
    from agentuniverse.base.config.application_configer.app_configer import AppConfiger
    from agentuniverse.base.config.application_configer.application_config_manager import ApplicationConfigManager

    managers = [ReaderManager(), StoreManager(), DocProcessorManager(), RagRouterManager()]
    old_pools = [dict(manager._instance_obj_map) for manager in managers]
    apps = ApplicationConfigManager()
    try:
        previous_app = apps.app_configer
    except ValueError:
        previous_app = None
    settings = Configer()
    settings.value = {"BASE_INFO": {"appname": "formula-roundtrip"}}
    apps.app_configer = AppConfiger().load_by_configer(settings)
    knowledge = None
    store = None
    try:
        for manager in managers:
            manager._instance_obj_map.clear()
        root = Path(__file__).resolve().parents[8]
        config = ComponentConfiger().load_by_configer(Configer(str(root / "agentuniverse/agent/action/knowledge/reader/file/xlsx_formula_reader.yaml")).load())
        assert config.get_component_config_type() == "READER"
        reader = XlsxFormulaReader().initialize_by_component_configer(config)
        keywords = JiebaKeywordExtractor(name="keywords")
        router = BaseRouter(name="base_router")
        store = SQLiteStore(name="local", db_path=str(tmp_path / "stored.db"), keyword_extractor="keywords")
        store._new_client()
        for manager, component in zip(managers, [reader, store, keywords, router]):
            manager.register(component.get_instance_code(), component)
        assert isinstance(ReaderManager().get_instance_obj("xlsx_formula_reader", strict=True), XlsxFormulaReader)
        source = tmp_path / "source.xlsx"
        source.write_bytes(workbook("5"))
        knowledge = Knowledge(stores=["local"], readers={"xlsx": "xlsx_formula_reader"})
        knowledge.insert_knowledge(source_path=str(source))
        assert store.conn.execute("SELECT COUNT(*) FROM documents").fetchone()[0] == 1
        store.conn.close()
        store._new_client()
        with patch.object(trace_module, "ConversationMemoryModule"), patch.object(trace_module, "Monitor") as monitor:
            monitor.get_invocation_chain.return_value = []
            result = knowledge.query_knowledge(query_str="alpha", keywords={"alpha"})
        assert len(result) == 1
        assert result[0].metadata["formula_cells"][0]["cached_value"] == "5"
        assert result[0].metadata["formula_cells"][0]["formula"] == "=A2+3"
        assert result[0].metadata["formula_results_verified"] is False
    finally:
        if knowledge:
            knowledge.insert_executor.shutdown(wait=True)
            knowledge.query_executor.shutdown(wait=True)
        if store:
            store.conn.close()
        for manager, previous in zip(managers, old_pools):
            manager._instance_obj_map.clear()
            manager._instance_obj_map.update(previous)
        apps.app_configer = previous_app
