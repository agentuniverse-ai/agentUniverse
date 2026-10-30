"""Exercise advertised XLSX stream inputs with real openpyxl workbooks."""

import io

import openpyxl
import pytest

from agentuniverse.agent.action.knowledge.reader.file.xlsx_reader import XlsxReader


@pytest.fixture
def workbook_bytes():
    """Build a workbook containing text, numeric, boolean and empty cells."""
    workbook = openpyxl.Workbook()
    first = workbook.active
    first.title = "First"
    first.append(["名称", "value"])
    first.append(["sample", 0])
    first.append([None, False])
    second = workbook.create_sheet("Second")
    second.append(["other", 3.5])
    workbook.create_sheet("Empty")
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


@pytest.mark.parametrize(
    "source_kind,expected_name",
    [
        ("unnamed", "unknown"),
        ("named_string", "source.xlsx"),
        ("named_path", "source.xlsx"),
        ("file_handle", "source.xlsx"),
        ("numeric_name", "unknown"),
        ("empty_name", "unknown"),
    ],
)
def test_binary_file_like_inputs(tmp_path, workbook_bytes, source_kind, expected_name):
    """Stream names become basenames; absent or non-path names use unknown."""
    if source_kind == "file_handle":
        path = tmp_path / "source.xlsx"
        path.write_bytes(workbook_bytes)
        source = path.open("rb")
    else:
        source = io.BytesIO(workbook_bytes)
        if source_kind == "named_string":
            source.name = str(tmp_path / "source.xlsx")
        elif source_kind == "named_path":
            source.name = tmp_path / "source.xlsx"
        elif source_kind == "numeric_name":
            source.name = 10
        elif source_kind == "empty_name":
            source.name = ""

    try:
        documents = XlsxReader().load_data(source)
        assert not source.closed
        assert [doc.metadata["sheet_name"] for doc in documents] == ["First", "Second"]
        assert all(doc.metadata["file_name"] == expected_name for doc in documents)
        assert documents[0].text == "名称 | value\nsample | 0\n | False"
        assert documents[1].text == "other | 3.5"
        assert documents[0].metadata["max_row"] == 3
        assert documents[0].metadata["max_col"] == 2
    finally:
        source.close()


@pytest.mark.parametrize("use_string_path", [False, True])
def test_existing_path_inputs(tmp_path, workbook_bytes, use_string_path):
    """Both public path input forms keep their existing basename metadata."""
    path = tmp_path / "source.xlsx"
    path.write_bytes(workbook_bytes)

    documents = XlsxReader().load_data(str(path) if use_string_path else path)

    assert len(documents) == 2
    assert all(doc.metadata["file_name"] == path.name for doc in documents)


def test_stream_ext_info_overrides_metadata(workbook_bytes):
    """Caller-provided metadata retains its existing override precedence."""
    with io.BytesIO(workbook_bytes) as source:
        documents = XlsxReader().load_data(
            source, ext_info={"file_name": "override.xlsx", "source": "upload"}
        )

    assert all(doc.metadata["file_name"] == "override.xlsx" for doc in documents)
    assert all(doc.metadata["source"] == "upload" for doc in documents)


def test_empty_workbook_stream():
    """An entirely empty workbook still produces no documents."""
    workbook = openpyxl.Workbook()
    with io.BytesIO() as source:
        workbook.save(source)
        workbook.close()
        source.seek(0)
        assert XlsxReader().load_data(source) == []
